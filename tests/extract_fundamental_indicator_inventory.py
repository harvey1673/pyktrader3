"""Extract workbook metadata rows into a normalized indicator catalog."""

import csv
import json
import os
import sys
from collections import Counter
from openpyxl import load_workbook


ROW_ALIASES = {
    "country": {"国家", "地区"},
    "indicator": {"指标名称"},
    "frequency": {"频率", "频度"},
    "unit": {"单位"},
    "indicator_id": {"指标ID", "指标编码"},
    "date_range": {"时间区间"},
    "source": {"来源", "数据来源"},
    "updated_at": {"更新时间"},
    "description": {"指标描述"},
}


def normalize(value):
    return "" if value is None else str(value).strip()


def extract_file(filename):
    workbook = load_workbook(filename, read_only=True, data_only=False)
    records = []
    sheet_counts = {}
    for sheet in workbook.worksheets:
        row_map = {}
        for row_idx in range(1, min(sheet.max_row, 25) + 1):
            label = normalize(sheet.cell(row_idx, 1).value)
            for field, aliases in ROW_ALIASES.items():
                if label in aliases:
                    row_map[field] = row_idx
        if "indicator" not in row_map:
            sheet_counts[sheet.title] = 0
            continue
        count = 0
        for col_idx in range(2, sheet.max_column + 1):
            indicator = normalize(sheet.cell(row_map["indicator"], col_idx).value)
            if not indicator:
                continue
            record = {
                "source_file": os.path.basename(filename),
                "sheet": sheet.title,
                "column": col_idx,
            }
            for field in ROW_ALIASES:
                row_idx = row_map.get(field)
                record[field] = normalize(sheet.cell(row_idx, col_idx).value) if row_idx else ""
            records.append(record)
            count += 1
        sheet_counts[sheet.title] = count
    return records, sheet_counts


def main():
    if len(sys.argv) < 3:
        raise SystemExit("usage: extract_fundamental_indicator_inventory.py OUTPUT.csv INPUT.xlsx ...")
    output = sys.argv[1]
    all_records = []
    summary = {}
    for filename in sys.argv[2:]:
        records, sheet_counts = extract_file(filename)
        all_records.extend(records)
        summary[os.path.basename(filename)] = sheet_counts
    fields = [
        "source_file", "sheet", "column", "country", "indicator", "frequency",
        "unit", "indicator_id", "date_range", "source", "updated_at", "description",
    ]
    with open(output, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(all_records)
    keywords = ["焦炭", "焦煤", "硅锰", "锰硅", "硅铁", "纯碱", "玻璃", "铜", "铝", "锌", "铅", "镍", "锡", "氧化铝"]
    keyword_counts = Counter()
    for record in all_records:
        for keyword in keywords:
            if keyword in record["indicator"]:
                keyword_counts[keyword] += 1
    print(json.dumps({"sheets": summary, "total_indicators": len(all_records), "keyword_counts": keyword_counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
