"""Extract formula-driven Level1/Level2 data requirements for secondary lead margin.

This script traces formulas in the secondary lead workbook and resolves upstream
lookup-backed indicators (name + ID) that are truly used by margin formulas.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Set, Tuple

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter


ROOT = Path("c:/dev/pyktrader3")
WORKBOOK = ROOT / "xldata/base/1_NEW secondary lead smelting margin.xlsx"
OUT_CSV = ROOT / "docs/secondary_lead_formula_requirements_level12.csv"
OUT_MD = ROOT / "docs/secondary_lead_formula_requirements_level12.md"


# Level definition is explicit to keep output stable and reviewable.
LEVEL_TARGETS: Dict[str, Dict[str, List[str]]] = {
    "Level1": {
        "economics-Henan": ["AU3", "BS3"],
        "economics-Anhui": ["AW3"],
    },
    "Level2": {
        "economics-Henan": ["BB3", "BD3"],
        "economics-Anhui": ["BD3", "BI3"],
    },
}


SHEET_CELL_REF_RE = re.compile(
    r"(?:'([^']+)'|([A-Za-z0-9_\-]+))!\$?([A-Z]{1,3})\$?(\d+)"
)
LOCAL_CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])\$?([A-Z]{1,3})\$?(\d+)")
VLOOKUP_RE = re.compile(
    r"VLOOKUP\(\s*[^,]+,\s*('[^']+'|[A-Za-z0-9_\-]+)!"
    r"\$?([A-Z]{1,3})\$?:\$?([A-Z]{1,3})\$?,\s*(\d+)",
    re.IGNORECASE,
)
VALID_ID_RE = re.compile(
    r"^(?:[sSmMgGlLaA]\d+|ID\d+|TODO_[A-Z0-9_]+|COMEX_[A-Z0-9_]+|"
    r"LBMA_[A-Z0-9_]+)$"
)


@dataclass(frozen=True)
class SourceItem:
    level: str
    target_sheet: str
    target_cell: str
    driver_sheet: str
    driver_cell: str
    source_sheet: str
    source_col: str
    source_name: str
    source_id: str


def normalize_sheet(sheet_token: str) -> str:
    """Normalize sheet names extracted from formulas."""
    return sheet_token.strip("'")


def parse_cell_refs(formula: str, current_sheet: str) -> List[Tuple[str, str]]:
    """Parse cell references and return (sheet, cell) pairs."""
    refs: List[Tuple[str, str]] = []
    for match in SHEET_CELL_REF_RE.finditer(formula):
        sheet_q, sheet_u, col, row = match.groups()
        sheet = normalize_sheet(sheet_q or sheet_u or current_sheet)
        refs.append((sheet, f"{col}{row}"))

    # Remove explicit sheet refs first to avoid duplicate local refs.
    local_formula = SHEET_CELL_REF_RE.sub("", formula)
    for match in LOCAL_CELL_REF_RE.finditer(local_formula):
        col, row = match.groups()
        refs.append((current_sheet, f"{col}{row}"))

    return refs


def parse_vlookups(formula: str) -> List[Tuple[str, str]]:
    """Parse VLOOKUP source mapping as (sheet, source_col)."""
    items: List[Tuple[str, str]] = []
    for match in VLOOKUP_RE.finditer(formula):
        raw_sheet, start_col, _, idx_str = match.groups()
        sheet = normalize_sheet(raw_sheet)
        idx = int(idx_str)
        start_idx = column_index_from_string(start_col)
        source_col = get_column_letter(start_idx + idx - 1)
        items.append((sheet, source_col))
    return items


def trace_sources(
    wb,
    level: str,
    target_sheet: str,
    target_cell: str,
) -> List[SourceItem]:
    """Trace upstream sources used by a target formula cell."""
    found: List[SourceItem] = []
    visited: Set[Tuple[str, str]] = set()

    def walk(sheet_name: str, cell_ref: str) -> None:
        key = (sheet_name, cell_ref)
        if key in visited:
            return
        visited.add(key)

        ws = wb[sheet_name]
        value = ws[cell_ref].value
        if not (isinstance(value, str) and value.startswith("=")):
            return

        for source_sheet, source_col in parse_vlookups(value):
            if source_sheet not in wb.sheetnames:
                continue
            source_ws = wb[source_sheet]
            source_name = source_ws[f"{source_col}1"].value
            source_id = source_ws[f"{source_col}2"].value
            if source_id in (None, "", "指标Id"):
                continue
            source_id_text = str(source_id)
            if not VALID_ID_RE.match(source_id_text):
                continue
            found.append(
                SourceItem(
                    level=level,
                    target_sheet=target_sheet,
                    target_cell=target_cell,
                    driver_sheet=sheet_name,
                    driver_cell=cell_ref,
                    source_sheet=source_sheet,
                    source_col=source_col,
                    source_name=str(source_name or ""),
                    source_id=source_id_text,
                )
            )

        for ref_sheet, ref_cell in parse_cell_refs(value, sheet_name):
            if ref_sheet in wb.sheetnames:
                walk(ref_sheet, ref_cell)

    walk(target_sheet, target_cell)
    return found


def write_outputs(items: List[SourceItem]) -> None:
    """Write CSV and Markdown outputs."""
    dedup: Dict[Tuple[str, str, str, str], SourceItem] = {}
    for item in items:
        key = (item.level, item.source_sheet, item.source_id, item.source_name)
        if key not in dedup:
            dedup[key] = item

    rows = sorted(
        dedup.values(),
        key=lambda x: (x.level, x.source_sheet, x.source_id, x.target_sheet),
    )

    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "level",
                "target_sheet",
                "target_cell",
                "source_sheet",
                "source_col",
                "source_name",
                "source_id",
            ],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "level": row.level,
                    "target_sheet": row.target_sheet,
                    "target_cell": row.target_cell,
                    "source_sheet": row.source_sheet,
                    "source_col": row.source_col,
                    "source_name": row.source_name,
                    "source_id": row.source_id,
                }
            )

    level_counts: Dict[str, int] = {}
    for row in rows:
        level_counts[row.level] = level_counts.get(row.level, 0) + 1

    with OUT_MD.open("w", encoding="utf-8") as f:
        f.write("# Secondary Lead Margin Formula Requirements (Level1/Level2)\n\n")
        f.write("来源: 通过追踪 Excel 公式得到，非人工主观筛选。\n\n")
        f.write(f"- 工作簿: {WORKBOOK.name}\n")
        f.write(f"- 输出CSV: {OUT_CSV.name}\n")
        f.write(f"- Level1 条数: {level_counts.get('Level1', 0)}\n")
        f.write(f"- Level2 条数: {level_counts.get('Level2', 0)}\n\n")
        f.write("## 目标公式\n\n")
        for level, by_sheet in LEVEL_TARGETS.items():
            f.write(f"- {level}:\n")
            for sheet, cells in by_sheet.items():
                f.write(f"  - {sheet}: {', '.join(cells)}\n")


def main() -> None:
    """Run extraction and write outputs."""
    wb = load_workbook(WORKBOOK, data_only=False)
    items: List[SourceItem] = []
    for level, by_sheet in LEVEL_TARGETS.items():
        for sheet, cells in by_sheet.items():
            for cell in cells:
                items.extend(trace_sources(wb, level, sheet, cell))

    write_outputs(items)
    print(f"generated {OUT_CSV}")
    print(f"generated {OUT_MD}")


if __name__ == "__main__":
    main()
