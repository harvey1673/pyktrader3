"""Audit inventory metadata and cached observations in the named workbooks."""

import csv
import datetime as dt
import json
import runpy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ALIASES = runpy.run_path(str(ROOT / "tests/extract_fundamental_indicator_inventory.py"))["ROW_ALIASES"]
INPUT = Path("C:/Users/harve/Nutstore/1/Nutstore")
OUTPUT = ROOT / "docs/commodity_inventory_research"


def text(value):
    if value is None:
        return ""
    return value.isoformat() if isinstance(value, (dt.date, dt.datetime)) else str(value).strip()


def classify(name, unit):
    if any(word in name for word in ("平均价", "价格", "报价")):
        return "非库存：价格"
    if "未执行合同" in name:
        return "非库存：订单"
    if any(word in name for word in ("PMI", "采购经理", "同比", "变化")):
        return "辅助：指数或变动量"
    if "存栏" in name:
        return "辅助：存栏"
    if "仓单预报" in name:
        return "辅助：仓单预报"
    if any(word in name for word in ("库存天数", "可用天数", "库存消费比")) or unit in ("天", "日"):
        return "库存天数"
    if "库容" in name:
        return "库容比"
    if "仓单" in name:
        return "仓单数量"
    return "库存数量"


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    records, summary = [], {}
    for name in ("ifind_daily.xlsx", "ifind_data.xlsx", "mysteel_metal.xlsx"):
        path = INPUT / name
        workbook = load_workbook(path, read_only=True, data_only=True)
        counts = {}
        for sheet in workbook:
            headers = list(sheet.iter_rows(max_row=min(25, sheet.max_row), values_only=True))
            row_map = {}
            for index, row in enumerate(headers):
                for field, aliases in ALIASES.items():
                    if row and text(row[0]) in aliases:
                        row_map[field] = index
            if "indicator" not in row_map:
                counts[sheet.title] = {"indicators": 0, "inventory": 0}
                continue
            selected = {}
            indicators = 0
            for col, value in enumerate(headers[row_map["indicator"]][1:], 1):
                if not text(value):
                    continue
                indicators += 1
                if not any(word in text(value) for word in ("库存", "仓单", "存栏", "库容", "存货", "结转", "储备")):
                    continue
                record = {"source_file": name, "sheet": sheet.title, "column": col + 1}
                for field in ALIASES:
                    index = row_map.get(field)
                    record[field] = text(headers[index][col]) if index is not None and col < len(headers[index]) else ""
                record["indicator_type"] = classify(record["indicator"], record["unit"])
                record.update(first_observation="", last_observation="", last_value="", observation_count=0)
                selected[col] = record
            if selected:
                for row in sheet.iter_rows(values_only=True):
                    if not row or not isinstance(row[0], (dt.date, dt.datetime)):
                        continue
                    for col, record in selected.items():
                        value = row[col] if col < len(row) else None
                        if not isinstance(value, (int, float)) or isinstance(value, bool):
                            continue
                        date = text(row[0])[:10]
                        record["observation_count"] += 1
                        if not record["first_observation"] or date < record["first_observation"]:
                            record["first_observation"] = date
                        if not record["last_observation"] or date >= record["last_observation"]:
                            record["last_observation"], record["last_value"] = date, text(value)
            records.extend(selected.values())
            counts[sheet.title] = {"indicators": indicators, "inventory": len(selected)}
        workbook.close()
        summary[name] = {"path": str(path), "modified_at": dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(), "sheets": counts}
    fields = list(records[0])
    with (OUTPUT / "existing_inventory_indicators.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fields)
        writer.writeheader()
        writer.writerows(records)
    (OUTPUT / "workbook_audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"workbooks": summary, "inventory_columns": len(records)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
