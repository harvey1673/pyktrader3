import datetime as dt
import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from misc_scripts.strategy_scenario_backtest import (
    _append_spread_contract_prices,
    _strategy_assets,
    FactorFrameSignalProvider,
    SignalBacktestResult,
    SignalSpec,
    StrategyScenario,
    btmetrics_result,
    compare_portfolios,
    compose_portfolio,
    load_excel_weight_scenarios,
    load_strategy_scenario,
    run_strategy_comparison,
    signal_asset_diagnostics,
    signal_coverage,
    write_comparison_excel,
    write_comparison_html,
    write_unit_signal_pnl_csv,
)


class StaticProvider:
    def __init__(self, results):
        self.results = results
        self.calls = []

    def backtest(self, spec):
        self.calls.append(spec.cache_key())
        return self.results[spec.factor_name]


class FakeMetricsBase:
    """Dependency-light implementation of the btmetrics methods used here."""

    def __init__(
        self,
        holdings,
        returns,
        shift_holdings=0,
        cost_dict=None,
        **_kwargs,
    ):
        index = holdings.index.intersection(returns.dropna(how="all").index)
        columns = holdings.columns.intersection(returns.columns)
        self.holdings = holdings.reindex(index=index, columns=columns).shift(
            shift_holdings
        )
        self.returns = returns.reindex(index=index, columns=columns)
        self.cost_dict = dict(cost_dict or {})

    def calculate_pnl_stats(self, **_kwargs):
        costs = self.holdings.diff().abs().multiply(pd.Series(self.cost_dict))
        return {"asset_pnl": self.holdings.multiply(self.returns) - costs}

    def calculate_daily_pnl(self, trade_prices, close_prices, mode="ret"):
        if mode != "ret":
            raise AssertionError("The scenario provider should use return PNL")
        pnl = self.holdings.multiply(close_prices.pct_change().fillna(0.0))
        pnl -= self.holdings.diff().abs().multiply(pd.Series(self.cost_dict))
        trades = self.holdings - self.holdings.shift(1).fillna(0.0)
        pnl += trades.multiply(close_prices / trade_prices - 1.0)
        return pnl


def _path(holdings, gross_pnl=None, rate=0.01, bucket="default"):
    dates = pd.date_range("2024-01-01", periods=len(holdings), freq="B")
    holding_df = pd.DataFrame({"cu": holdings}, index=dates, dtype=float)
    if gross_pnl is None:
        gross_pnl = np.zeros(len(holdings))
    pnl_df = pd.DataFrame({"cu": gross_pnl}, index=dates, dtype=float)
    return SignalBacktestResult(
        holdings=holding_df,
        gross_asset_pnl=pnl_df,
        cost_rates=pd.Series({"cu": rate}),
        execution_bucket=bucket,
    )


def _scenario(name, specs, scaler=1.0):
    return StrategyScenario(
        name=name,
        strategy_file="sample.json",
        scaler=scaler,
        signals={spec.factor_name: spec for spec in specs},
        config={},
    )


class StrategyScenarioBacktestTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_strategy_assets_include_every_json_underlier(self):
        misc_module = types.ModuleType("pycmqlib3.utility.misc")
        misc_module.inst2product = lambda value: value.rstrip("0123456789")
        strategy = StrategyScenario(
            name="test",
            strategy_file="test.json",
            scaler=1.0,
            signals={},
            config={
                "assets": [
                    {"underliers": ["rb2610", "hc2610"]},
                    {"underliers": ["rb2701"]},
                ]
            },
        )

        with patch.dict(sys.modules, {"pycmqlib3.utility.misc": misc_module}):
            self.assertEqual(_strategy_assets(strategy), ["rb", "hc"])

    def test_spread_price_loader_builds_notebook_d1_columns(self):
        dates = pd.date_range("2015-01-02", periods=4, freq="B")
        futures = pd.DataFrame(
            {
                ("rbc1", "close"): [100.0, 101.0, 102.0, 103.0],
                ("hcc1", "close"): [90.0, 91.0, 92.0, 93.0],
            },
            index=dates,
        )
        calls = []

        def loader(asset, **kwargs):
            calls.append((asset, kwargs))
            offset = 10.0 if asset == "rb" else 20.0
            return pd.DataFrame(
                {"close": np.arange(4.0) + offset}, index=dates
            )

        output = _append_spread_contract_prices(
            futures,
            [SignalSpec("factor", "spread_signal", "pos", 1.0)],
            ["rb", "hc"],
            factors_by_spread2={"spread_signal": ["spd_hc_rb_c1"]},
            spread_config={
                "spd_hc_rb_c1": [[("hc", 1.0), ("rb", -1.0)], 20, "-30b", 1]
            },
            end_date=dates[-1].date(),
            price_loader=loader,
        )

        self.assertIn(("rbd1", "close"), output.columns)
        self.assertIn(("hcd1", "close"), output.columns)
        self.assertEqual({asset for asset, _kwargs in calls}, {"rb", "hc"})
        for _asset, kwargs in calls:
            self.assertEqual(kwargs["roll_rule"], "-30b")
            self.assertEqual(kwargs["shift_mode"], 1)
            self.assertEqual(kwargs["start_date"], dt.date(2015, 1, 2))

    def _write_strategy(self):
        settings = self.root / "settings"
        settings.mkdir()
        data = {
            "class": "example.Strategy",
            "config": {
                "name": "sample",
                "pos_scaler": 10,
                "factor_repo": {
                    "factor.one": {
                        "name": "signal_one",
                        "type": "ts",
                        "exec_assets": ["rb"],
                        "threshold": 0.25,
                        "rebal": 5,
                        "param": [1.0, 2.0],
                        "weight": 1.0,
                    }
                },
            },
        }
        (settings / "sample.json").write_text(
            json.dumps(data, indent=4), encoding="utf-8"
        )
        return settings

    def test_load_current_and_proposed_weights_from_excel_in_memory(self):
        settings = self._write_strategy()
        baseline = load_strategy_scenario(settings, "sample.json")
        workbook_path = self.root / "proposal.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "signal_weights"
        sheet.append(
            [
                "strategy",
                "factor_name",
                "signal_name",
                "type",
                "curr_weight",
                "new_weight",
            ]
        )
        sheet.append(
            ["sample.json", "factor.one", "signal_one_new", "xs-demean", 0.75, 0]
        )
        sheet.append(
            ["sample.json", "factor.two", "signal_two", "pos", 0.25, 2.5]
        )
        workbook.save(workbook_path)
        workbook.close()

        current, proposed = load_excel_weight_scenarios(baseline, workbook_path)

        self.assertEqual(baseline.signals["factor.one"].weight, 1.0)
        self.assertEqual(current.signals["factor.one"].weight, 0.75)
        self.assertEqual(current.signals["factor.two"].weight, 0.25)
        changed = proposed.signals["factor.one"]
        self.assertEqual(changed.weight, 0.0)
        self.assertEqual(changed.name, "signal_one_new")
        self.assertEqual(changed.exec_assets, ("rb",))
        self.assertEqual(changed.threshold, 0.25)
        self.assertEqual(changed.rebal, 5)
        self.assertEqual(changed.param, (1.0, 2.0))
        added = proposed.signals["factor.two"]
        self.assertEqual(added.exec_assets, ())
        self.assertEqual(added.threshold, 0.0)
        self.assertEqual(added.rebal, 1)
        self.assertEqual(added.param, (0.0, 0.0))
        self.assertEqual(added.weight, 2.5)

        saved = json.loads((settings / "sample.json").read_text(encoding="utf-8"))
        self.assertEqual(saved["config"]["factor_repo"]["factor.one"]["weight"], 1.0)

    def test_excel_scenario_loader_requires_both_weight_columns(self):
        settings = self._write_strategy()
        template = load_strategy_scenario(settings, "sample.json")
        workbook_path = self.root / "legacy.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "signal_weights"
        sheet.append(["strategy", "factor_name", "signal_name", "type", "weight"])
        sheet.append(["sample.json", "factor.one", "signal_one", "ts", 1.0])
        workbook.save(workbook_path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "curr_weight, new_weight"):
            load_excel_weight_scenarios(template, workbook_path)

    def test_identical_scenarios_have_zero_delta_and_reconcile(self):
        spec = SignalSpec("factor.one", "signal_one", "ts", 1.5)
        provider = StaticProvider(
            {"factor.one": _path([0, 1, 1, 0], [0, 0.02, -0.01, 0.01], rate=0.001)}
        )
        current = _scenario("current", [spec], scaler=2.0)
        proposed = _scenario("proposed", [spec], scaler=2.0)

        comparison = run_strategy_comparison(current, proposed, provider)

        self.assertTrue(np.allclose(comparison.daily_pnl["delta"], 0.0))
        standard_metrics = [
            metric
            for metric in comparison.summary.index
            if metric not in {"pnl_correlation", "information_ratio"}
        ]
        self.assertTrue(
            np.allclose(
                comparison.summary.loc[standard_metrics, "delta"].dropna(), 0.0
            )
        )
        self.assertAlmostEqual(comparison.summary.loc["pnl_correlation", "delta"], 1.0)
        self.assertTrue((comparison.checks["status"] == "OK").all())
        self.assertEqual(
            list(comparison.current_btmetrics.portfolio.columns), ["sharpe", "std"]
        )
        self.assertEqual(
            list(comparison.current_btmetrics.assets.columns),
            ["turnover", "pnl_per_trade"],
        )
        pd.testing.assert_frame_equal(
            comparison.current_btmetrics.portfolio,
            comparison.proposed_btmetrics.portfolio,
        )

    def test_btmetrics_result_matches_notebook_asset_formulas(self):
        spec = SignalSpec("factor.one", "signal_one", "ts", 1.0)
        provider = StaticProvider(
            {"factor.one": _path([0, 1, 1, 0], [0, 0.02, -0.01, 0.01], rate=0.001)}
        )
        portfolio = compose_portfolio(_scenario("current", [spec]), provider)

        stats = btmetrics_result(portfolio, tenors=["3m"])

        expected_turnover = (
            100.0
            * portfolio.aggregate_holdings.diff().abs().mean()["cu"]
            / portfolio.aggregate_holdings.abs().mean()["cu"]
        )
        expected_pnl_per_trade = (
            10000.0
            * portfolio.net_asset_pnl.mean()["cu"]
            / portfolio.aggregate_holdings.diff().abs().mean()["cu"]
        )
        self.assertAlmostEqual(stats.assets.loc["cu", "turnover"], expected_turnover)
        self.assertAlmostEqual(
            stats.assets.loc["cu", "pnl_per_trade"], expected_pnl_per_trade
        )
        self.assertEqual(list(stats.portfolio.index), ["full", "3m"])

    def test_signal_coverage_reports_first_active_date(self):
        spec = SignalSpec("factor.one", "signal_one", "ts", 1.0)
        provider = StaticProvider(
            {"factor.one": _path([0, 1, 1, 0], [0, 0.02, -0.01, 0.01], rate=0.0)}
        )
        comparison = run_strategy_comparison(
            _scenario("current", [spec]),
            _scenario("proposed", [spec]),
            provider,
        )

        coverage = signal_coverage(comparison)

        self.assertEqual(
            coverage.loc["factor.one", "first_active_current"],
            pd.Timestamp("2024-01-02"),
        )
        self.assertEqual(coverage.loc["factor.one", "active_days_current"], 3)

    def test_signal_asset_diagnostics_reconcile_to_signal_pnl(self):
        spec = SignalSpec("factor.one", "signal_one", "ts", 1.0)
        portfolio = compose_portfolio(
            _scenario("current", [spec]),
            StaticProvider(
                {"factor.one": _path([0, 1, 1, 0], [0, 0.02, -0.01, 0.01], rate=0.0)}
            ),
        )

        diagnostics = signal_asset_diagnostics(portfolio)

        pd.testing.assert_series_equal(
            portfolio.signal_asset_pnl["factor.one"].sum(axis=1),
            portfolio.signal_pnl["factor.one"],
            check_names=False,
        )
        self.assertIn(("factor.one", "Total"), diagnostics.index)
        self.assertIn(("factor.one", "cu"), diagnostics.index)
        self.assertAlmostEqual(
            diagnostics.loc[("factor.one", "Total"), "daily_std"],
            portfolio.signal_pnl["factor.one"].std(),
        )

    def test_unit_signal_pnl_excludes_weight_and_scaler(self):
        spec = SignalSpec("factor.one", "signal_one", "ts", 2.0)
        portfolio = compose_portfolio(
            _scenario("current", [spec], scaler=10.0),
            StaticProvider(
                {"factor.one": _path([0, 1, 1, 0], [0, 0.02, 0.01, -0.01], rate=0.001)}
            ),
        )

        expected = pd.Series(
            [0.0, 0.019, 0.01, -0.011],
            index=portfolio.unit_signal_pnl.index,
            name="factor.one",
        )
        pd.testing.assert_series_equal(
            portfolio.unit_signal_pnl["factor.one"], expected
        )
        self.assertFalse(
            np.allclose(
                portfolio.unit_signal_pnl["factor.one"],
                portfolio.signal_pnl["factor.one"],
            )
        )

    def test_comparison_saves_zero_weight_signal_for_optimization(self):
        active = SignalSpec("active", "active", "ts", 1.0)
        candidate = SignalSpec("candidate", "candidate", "ts", 0.0)
        scenario = _scenario("current", [active, candidate])
        provider = StaticProvider(
            {
                "active": _path([0, 1, 1], [0, 0.01, 0.01], rate=0.0),
                "candidate": _path([0, -1, -1], [0, 0.02, -0.01], rate=0.0),
            }
        )

        comparison = run_strategy_comparison(scenario, scenario, provider)

        self.assertEqual(
            list(comparison.proposed.unit_signal_pnl.columns),
            ["active", "candidate"],
        )
        self.assertEqual(list(comparison.proposed.signal_pnl.columns), ["active"])

    def test_netted_costs_recognize_offsetting_trades(self):
        spec_a = SignalSpec("a", "a", "ts", 1.0)
        spec_b = SignalSpec("b", "b", "ts", 1.0)
        scenario = _scenario("portfolio", [spec_a, spec_b])
        provider = StaticProvider(
            {
                "a": _path([0, 1, 0], rate=0.01),
                "b": _path([0, -1, 0], rate=0.01),
            }
        )

        sleeve = compose_portfolio(scenario, provider, cost_mode="sleeve")
        netted = compose_portfolio(scenario, provider, cost_mode="netted")

        self.assertGreater(float(sleeve.costs_by_asset.sum().sum()), 0.0)
        self.assertAlmostEqual(float(netted.costs_by_asset.sum().sum()), 0.0)
        self.assertAlmostEqual(float(netted.trade_volume.sum().sum()), 0.0)
        self.assertTrue(
            np.allclose(netted.signal_pnl.sum(axis=1), netted.portfolio_pnl)
        )

    def test_factor_frame_provider_applies_exclusions_and_lag(self):
        dates = pd.date_range("2024-01-01", periods=5, freq="B")
        factors = {
            "signal": pd.DataFrame(
                {"cu": [1, 2, 3, 4, 5], "rb": [5, 4, 3, 2, 1]}, index=dates
            )
        }
        returns = pd.DataFrame({"cu": 0.01, "rb": 0.02}, index=dates)
        volatility = pd.DataFrame({"cu": 0.1, "rb": 0.2}, index=dates)
        provider = FactorFrameSignalProvider(
            factors,
            returns,
            volatility,
            {"cu": 0.001, "rb": 0.001},
            holding_lag=1,
            metrics_class=FakeMetricsBase,
        )
        spec = SignalSpec(
            "factor", "signal", "ts", 1.0, exec_assets=("rb",), rebal=1
        )

        result = provider.backtest(spec)

        self.assertEqual(list(result.holdings.columns), ["cu"])
        self.assertEqual(result.holdings.iloc[0, 0], 0.0)
        self.assertEqual(result.holdings.iloc[1, 0], 10.0)
        self.assertAlmostEqual(result.gross_asset_pnl.iloc[2, 0], 0.2)

    def test_factor_provider_uses_configured_point_in_time_execution_price(self):
        dates = pd.date_range("2024-01-01", periods=5, freq="B")
        factors = {
            "auag_cme_wratio_zs": pd.DataFrame(
                {"cu": [1.0, 2.0, 3.0, 4.0, 5.0]}, index=dates
            )
        }
        close = pd.DataFrame(
            {"cu": [100.0, 100.0, 100.0, 110.0, 110.0]}, index=dates
        )
        a1505 = pd.DataFrame(
            {"cu": [100.0, 100.0, 100.0, 105.0, 110.0]}, index=dates
        )
        n305 = pd.DataFrame(
            {"cu": [100.0, 100.0, 100.0, 108.0, 110.0]}, index=dates
        )
        provider = FactorFrameSignalProvider(
            factors,
            close.pct_change(),
            pd.DataFrame({"cu": 1.0}, index=dates),
            {"cu": 0.001},
            close_prices=close,
            execution_prices={"a1505": a1505, "n305": n305},
            execution_config={
                "auag_cme_wratio_zs": {"win": "a1505", "lag": 1}
            },
            default_execution_window="n305",
            metrics_class=FakeMetricsBase,
        )

        result = provider.backtest(
            SignalSpec("factor", "auag_cme_wratio_zs", "pos", 1.0)
        )

        self.assertEqual(result.execution_bucket, "a1505")
        expected = 0.20 + (110.0 / 105.0 - 1.0)
        self.assertAlmostEqual(result.gross_asset_pnl.iloc[2, 0], expected)
        wrong_night_price_pnl = 0.20 + (110.0 / 108.0 - 1.0)
        self.assertNotAlmostEqual(
            result.gross_asset_pnl.iloc[2, 0], wrong_night_price_pnl
        )

    def test_factor_provider_prefers_json_factor_key_for_shared_signal_name(self):
        dates = pd.date_range("2024-01-01", periods=4, freq="B")
        provider = FactorFrameSignalProvider(
            {
                "shared.ts": pd.DataFrame({"cu": 1.0}, index=dates),
                "shared.xs": pd.DataFrame({"cu": 3.0}, index=dates),
            },
            pd.DataFrame({"cu": 0.01}, index=dates),
            pd.DataFrame({"cu": 1.0}, index=dates),
            {"cu": 0.0},
            holding_lag=1,
            metrics_class=FakeMetricsBase,
        )

        ts_result = provider.backtest(SignalSpec("shared.ts", "shared", "ts", 1.0))
        xs_result = provider.backtest(SignalSpec("shared.xs", "shared", "ts", 1.0))

        self.assertAlmostEqual(float(ts_result.holdings.max().iloc[0]), 1.0)
        self.assertAlmostEqual(float(xs_result.holdings.max().iloc[0]), 3.0)

    def test_factor_provider_uses_spread_price_changes_and_volatility_override(self):
        dates = pd.date_range("2024-01-01", periods=6, freq="B")
        factors = {"spread": pd.DataFrame({"cu": 1.0}, index=dates)}
        close = pd.DataFrame({"cu": np.arange(100.0, 106.0)}, index=dates)
        spread_contract = pd.DataFrame(
            {"cu": np.arange(200.0, 212.0, 2.0)}, index=dates
        )
        provider = FactorFrameSignalProvider(
            factors,
            close.pct_change(),
            pd.DataFrame({"cu": 1.0}, index=dates),
            {"cu": 0.0},
            close_prices=close,
            execution_prices={"close": close},
            execution_config={"spread": {"win": "close", "lag": 1}},
            volatility_overrides={
                "spread": pd.DataFrame({"cu": 2.0}, index=dates)
            },
            traded_price_overrides={"spread": spread_contract},
            pnl_modes={"spread": "px"},
            metrics_class=FakeMetricsBase,
        )

        result = provider.backtest(SignalSpec("factor", "spread", "ts", 1.0))

        self.assertAlmostEqual(float(result.holdings.max().iloc[0]), 0.5)
        self.assertAlmostEqual(float(result.gross_asset_pnl.max().iloc[0]), 1.0)

    def test_factor_provider_applies_signal_store_buffer_post_function(self):
        dates = pd.date_range("2024-01-01", periods=5, freq="B")
        seen = []

        def fake_buffer(frame, size):
            seen.append(size)
            return frame * 0.25

        provider = FactorFrameSignalProvider(
            {"signal": pd.DataFrame({"cu": 1.0}, index=dates)},
            pd.DataFrame({"cu": 0.01}, index=dates),
            pd.DataFrame({"cu": 1.0}, index=dates),
            {"cu": 0.0},
            post_funcs={"signal": "raw|buf0.4"},
            signal_buffer_func=fake_buffer,
            metrics_class=FakeMetricsBase,
        )

        result = provider.backtest(SignalSpec("factor", "signal", "ts", 1.0))

        self.assertEqual(seen, [0.4])
        self.assertAlmostEqual(float(result.holdings.max().iloc[0]), 0.25)

    def test_xs_type_uses_execution_config_suffix(self):
        dates = pd.date_range("2024-01-01", periods=4, freq="B")
        factor = pd.DataFrame(
            {"au": [1.0, 2.0, 3.0, 4.0], "ag": [-1.0, -2.0, -3.0, -4.0]},
            index=dates,
        )
        prices = pd.DataFrame(
            {"au": [100.0, 101.0, 102.0, 103.0], "ag": [100.0, 99.0, 98.0, 97.0]},
            index=dates,
        )
        provider = FactorFrameSignalProvider(
            {"auag_etf_mrev": factor},
            prices.pct_change(),
            pd.DataFrame(1.0, index=dates, columns=["au", "ag"]),
            {"au": 0.0, "ag": 0.0},
            close_prices=prices,
            execution_prices={"a1505": prices},
            execution_config={
                "auag_etf_mrev_xdemean": {"win": "a1505", "lag": 1}
            },
            default_execution_window="n305",
            metrics_class=FakeMetricsBase,
        )

        result = provider.backtest(
            SignalSpec("auag_etf_mrev.xs", "auag_etf_mrev", "xs-demean", 1.0)
        )

        self.assertEqual(result.execution_bucket, "a1505")

    def test_unsupported_factor_fails_fast(self):
        dates = pd.date_range("2024-01-01", periods=3, freq="B")
        provider = FactorFrameSignalProvider(
            {},
            pd.DataFrame({"cu": 0.0}, index=dates),
            pd.DataFrame({"cu": 0.1}, index=dates),
            {"cu": 0.001},
            metrics_class=FakeMetricsBase,
        )
        spec = SignalSpec("missing", "not_in_store", "ts", 1.0)
        with self.assertRaisesRegex(KeyError, "Historical factor data is unavailable"):
            provider.backtest(spec)

    def test_write_comparison_excel(self):
        old_spec = SignalSpec("factor.one", "signal_one", "ts", 1.0)
        new_spec = SignalSpec("factor.one", "signal_one", "ts", 2.0)
        provider = StaticProvider(
            {"factor.one": _path([0, 1, 1, 0], [0, 0.02, 0.01, -0.01], rate=0.001)}
        )
        comparison = run_strategy_comparison(
            _scenario("current", [old_spec]),
            _scenario("proposed", [new_spec]),
            provider,
        )
        output = self.root / "comparison.xlsx"

        result = write_comparison_excel(comparison, output)

        self.assertEqual(result, output.resolve())
        workbook = load_workbook(output, data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "Current Portfolio",
                "Proposed Portfolio",
                "Current Assets",
                "Proposed Assets",
                "Asset Comparison",
                "Signal Attribution",
                "Current Signal PNL",
                "Proposed Signal PNL",
                "Current Signal Contribution",
                "Proposed Signal Contribution",
                "Current Signal Metrics",
                "Proposed Signal Metrics",
                "Signal Coverage",
                "Daily PNL",
                "Checks",
                "Run Info",
            ],
        )
        self.assertEqual(len(workbook["Summary"]._charts), 1)
        self.assertEqual(workbook["Checks"]["F2"].value, "OK")
        workbook.close()

        csv_output = self.root / "comparison_signal_pnl.csv"
        self.assertEqual(
            write_unit_signal_pnl_csv(comparison, csv_output),
            csv_output.resolve(),
        )
        saved_pnl = pd.read_csv(csv_output, index_col="date", parse_dates=["date"])
        pd.testing.assert_frame_equal(
            saved_pnl,
            comparison.proposed.unit_signal_pnl.rename_axis("date"),
            check_freq=False,
        )

    def test_write_comparison_html(self):
        old_spec = SignalSpec("factor.one", "signal_one", "ts", 1.0)
        new_spec = SignalSpec("factor.one", "signal_one", "ts", 2.0)
        provider = StaticProvider(
            {"factor.one": _path([0, 1, 1, 0], [0, 0.02, 0.01, -0.01], rate=0.001)}
        )
        comparison = run_strategy_comparison(
            _scenario("current", [old_spec]),
            _scenario("proposed", [new_spec]),
            provider,
        )
        output = self.root / "comparison.html"

        with (
            patch(
                "misc_scripts.strategy_scenario_backtest._comparison_chart_html",
                return_value=(
                    '<div id="cumulative-chart"></div>',
                    '<div id="drawdown-chart"></div>',
                    '<div id="distribution-chart"></div>',
                ),
            ),
            patch(
                "misc_scripts.strategy_scenario_backtest._signal_chart_html",
                return_value=(
                    '<div id="signal-path-chart"></div>',
                    '<div id="signal-contribution-chart"></div>',
                ),
            ),
            patch(
                "misc_scripts.strategy_scenario_backtest._signal_asset_chart_html",
                side_effect=(
                    '<div id="current-signal-asset-chart"></div>',
                    '<div id="proposed-signal-asset-chart"></div>',
                ),
            ),
        ):
            result = write_comparison_html(comparison, output)

        document = output.read_text(encoding="utf-8")
        self.assertEqual(result, output.resolve())
        self.assertIn("<!doctype html>", document)
        self.assertIn("Portfolio Backtest Comparison", document)
        self.assertIn("Executive Summary", document)
        self.assertIn("Portfolio performance table", document)
        self.assertIn("Sharpe ratio", document)
        self.assertIn("Portfolio btmetrics by tenor", document)
        self.assertIn("Asset trading efficiency", document)
        self.assertIn("PNL per trade", document)
        self.assertIn("cumulative-chart", document)
        self.assertIn("drawdown-chart", document)
        self.assertIn("distribution-chart", document)
        self.assertNotIn("portfolio-btmetrics-chart", document)
        self.assertNotIn("asset-btmetrics-chart", document)
        self.assertIn("signal-path-chart", document)
        self.assertIn("signal-contribution-chart", document)
        self.assertIn("current-signal-asset-chart", document)
        self.assertIn("proposed-signal-asset-chart", document)
        self.assertIn("Signal diagnostics", document)
        self.assertIn("Signal data coverage", document)
        self.assertIn("signal_execution_config", document)
        self.assertIn("a1505", document)
        self.assertIn("n305", document)


if __name__ == "__main__":
    unittest.main()
