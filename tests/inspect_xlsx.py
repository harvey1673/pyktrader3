import json
import os
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

for filename in sys.argv[1:]:
    print(f"\n### {os.path.basename(filename)}")
    wb = load_workbook(filename, read_only=True, data_only=False)
    for ws in wb.worksheets:
        print(f"\n## {ws.title} rows={ws.max_row} cols={ws.max_column}")
        shown = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True):
            values = list(row[:min(ws.max_column, 60)])
            if any(v not in (None, "") for v in values):
                print(json.dumps(values, ensure_ascii=False, default=str))
                shown += 1
                if shown >= 12:
                    break
