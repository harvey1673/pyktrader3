import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from misc_scripts.signal_weight_excel import (
    WORKBOOK_COLUMNS,
    export_signal_weights_to_excel,
    generate_strategy_json_from_excel,
    import_signal_weights_from_excel,
)


def _write_strategy(path: Path) -> None:
    data = {
        "class": "example.Strategy",
        "config": {
            "factor_repo": {
                "existing.factor": {
                    "name": "old_name",
                    "type": "ts",
                    "exec_assets": ["cu"],
                    "threshold": 0.25,
                    "rebal": 5,
                    "param": [1.0, 2.0],
                    "weight": 1.5,
                    "custom_field": {"keep": True},
                }
            },
            "unrelated": "preserve me",
        },
    }
    path.write_text(json.dumps(data, indent=4), encoding="utf-8")


class SignalWeightExcelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.settings_dir = self.root / "settings"
        self.settings_dir.mkdir()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_export_signal_weights_to_excel(self) -> None:
        _write_strategy(self.settings_dir / "sample.json")
        excel_path = self.root / "weights.xlsx"

        result = export_signal_weights_to_excel(self.settings_dir, excel_path)

        self.assertEqual(result, excel_path.resolve())
        workbook = load_workbook(excel_path, data_only=True)
        worksheet = workbook["signal_weights"]
        self.assertEqual(
            tuple(cell.value for cell in worksheet[1]), WORKBOOK_COLUMNS
        )
        self.assertEqual(
            tuple(cell.value for cell in worksheet[2]),
            (
                "sample.json",
                "existing.factor",
                "old_name",
                "ts",
                1.5,
                1.5,
            ),
        )
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(list(worksheet.tables), ["SignalWeights"])
        self.assertEqual(worksheet["F2"].font.color.rgb, "000000FF")
        workbook.close()

    def test_import_preserves_existing_fields_and_adds_defaults(self) -> None:
        strategy_path = self.settings_dir / "sample.json"
        _write_strategy(strategy_path)

        excel_path = self.root / "weights.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "signal_weights"
        worksheet.append(WORKBOOK_COLUMNS)
        worksheet.append(
            ["sample.json", "existing.factor", "new_name", "xs", 1.5, 2.75]
        )
        worksheet.append(["sample", "new.factor", "brand_new", "pos", 0.0, -0.5])
        worksheet.append(["missing.json", None, None, None, None, None])
        workbook.save(excel_path)
        workbook.close()

        result = import_signal_weights_from_excel(self.settings_dir, excel_path)

        self.assertEqual(result.processed_rows, 3)
        self.assertEqual(result.updated_factors, 1)
        self.assertEqual(result.added_factors, 1)
        self.assertEqual(result.skipped_missing_strategies, 1)
        self.assertEqual(result.written_files, (strategy_path.resolve(),))

        data = json.loads(strategy_path.read_text(encoding="utf-8"))
        repo = data["config"]["factor_repo"]
        existing = repo["existing.factor"]
        self.assertEqual(existing["name"], "new_name")
        self.assertEqual(existing["type"], "xs")
        self.assertEqual(existing["weight"], 2.75)
        self.assertEqual(existing["exec_assets"], ["cu"])
        self.assertEqual(existing["threshold"], 0.25)
        self.assertEqual(existing["rebal"], 5)
        self.assertEqual(existing["param"], [1.0, 2.0])
        self.assertEqual(existing["custom_field"], {"keep": True})
        self.assertEqual(data["config"]["unrelated"], "preserve me")

        self.assertEqual(
            repo["new.factor"],
            {
                "name": "brand_new",
                "type": "pos",
                "exec_assets": [],
                "threshold": 0.0,
                "rebal": 1,
                "param": [0.0, 0.0],
                "weight": -0.5,
            },
        )

    def test_generate_proposed_json_without_changing_source(self) -> None:
        strategy_path = self.settings_dir / "sample.json"
        _write_strategy(strategy_path)
        source = json.loads(strategy_path.read_text(encoding="utf-8"))
        source["config"]["factor_repo"]["removed.factor"] = {
            "name": "removed",
            "type": "ts",
            "weight": 1.0,
        }
        strategy_path.write_text(json.dumps(source, indent=4), encoding="utf-8")
        original = strategy_path.read_text(encoding="utf-8")

        excel_path = self.root / "weights.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "signal_weights"
        worksheet.append(WORKBOOK_COLUMNS)
        worksheet.append(
            ["sample.json", "existing.factor", "new_name", "xs", 1.5, 2.75]
        )
        worksheet.append(
            ["sample", "new.factor", "brand_new", "pos", 0.0, -0.5]
        )
        worksheet.append(
            ["sample", "removed.factor", "removed", "ts", 1.0, 0.0]
        )
        worksheet.append(
            ["sample", "zero.new.factor", "zero_new", "ts", 0.0, 0.0]
        )
        worksheet.append(
            ["other.json", "ignored.factor", "ignored", "pos", 1.0, 9.0]
        )
        workbook.save(excel_path)
        workbook.close()

        result = generate_strategy_json_from_excel(strategy_path, excel_path)

        expected_path = self.root / "sample_proposed.json"
        self.assertEqual(result, expected_path.resolve())
        self.assertEqual(strategy_path.read_text(encoding="utf-8"), original)
        data = json.loads(expected_path.read_text(encoding="utf-8"))
        repo = data["config"]["factor_repo"]
        self.assertEqual(set(repo), {"existing.factor", "new.factor"})
        self.assertEqual(repo["existing.factor"]["weight"], 2.75)
        self.assertEqual(repo["existing.factor"]["name"], "new_name")
        self.assertEqual(repo["existing.factor"]["type"], "xs")
        self.assertEqual(repo["existing.factor"]["custom_field"], {"keep": True})
        self.assertEqual(
            repo["new.factor"],
            {
                "name": "brand_new",
                "type": "pos",
                "exec_assets": [],
                "threshold": 0.0,
                "rebal": 1,
                "param": [0.0, 0.0],
                "weight": -0.5,
            },
        )
        self.assertEqual(data["config"]["unrelated"], "preserve me")


if __name__ == "__main__":
    unittest.main()
