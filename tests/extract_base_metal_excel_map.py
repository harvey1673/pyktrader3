"""Extract code/ticker candidates from base-metal Excel workbooks.

This helper scans the first rows of all workbooks under ``xldata/base`` and
creates an inventory CSV in ``docs/`` to support margin-model data mapping.
"""

from __future__ import annotations

import csv
import importlib.util
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BASE_XLSX_DIR = ROOT / "xldata" / "base"
INDEX_MAP_FILE = ROOT / "tests" / "index_map_full.py"
OUTPUT_FILE = ROOT / "docs" / "base_metal_excel_code_inventory.csv"

# iFind-like codes, SMM ids, and Reuters tickers commonly seen in these files.
IFIND_CODE_RE = re.compile(r"^[SMLG]\d{6,}(?:\.\d+)?$")
SMM_ID_RE = re.compile(r"^s\d{8}$", re.IGNORECASE)
REUTERS_TICKER_RE = re.compile(r"^[A-Za-z]{2,6}[A-Za-z0-9]{0,4}=?$")


@dataclass(frozen=True)
class ExtractedCode:
    workbook: str
    sheet: str
    code_type: str
    raw_code: str
    mapped_alias: str
    context_text: str


def load_index_map(index_map_path: Path) -> Dict[str, str]:
    """Load index_map_full from a Python module path."""
    spec = importlib.util.spec_from_file_location("index_map_full_module", index_map_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load module: {index_map_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return dict(module.index_map_full)


def classify_token(token: str) -> str | None:
    """Classify token as ifind/smm/reuters candidate or return None."""
    token = token.strip()
    if IFIND_CODE_RE.match(token):
        return "ifind_code"
    if SMM_ID_RE.match(token):
        return "smm_id"
    if REUTERS_TICKER_RE.match(token):
        banned = {
            "DATE",
            "WIND",
            "SMM",
            "VAT",
            "COST",
            "PRICE",
            "LME",
            "RMB",
            "USD",
            "YEAR",
            "MONTH",
            "WEEK",
        }
        if token.upper() in banned:
            return None
        # Reuters-like symbols are usually contract-like tokens,
        # often with digits or trailing '=' for FX series.
        if not any(ch.isdigit() for ch in token) and not token.endswith("="):
            return None
        return "reuters_ticker"
    return None


def tokenize_cell_text(value: str) -> Iterable[str]:
    """Yield code-like tokens from a single cell string."""
    text = value.strip()
    if not text:
        return []
    return re.findall(r"[A-Za-z0-9.=]+", text)


def scan_workbook(path: Path, index_map: Dict[str, str]) -> List[ExtractedCode]:
    """Scan a workbook for code candidates in the first 80 rows of each sheet."""
    extracted: List[ExtractedCode] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    seen: set[tuple[str, str, str]] = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=80, values_only=True):
            for cell in row:
                if cell is None:
                    continue
                text = str(cell).strip()
                if not text:
                    continue
                for token in tokenize_cell_text(text):
                    code_type = classify_token(token)
                    if code_type is None:
                        continue
                    normalized = token.strip()
                    key = (sheet_name, code_type, normalized)
                    if key in seen:
                        continue
                    seen.add(key)
                    mapped = index_map.get(normalized, "")
                    extracted.append(
                        ExtractedCode(
                            workbook=path.name,
                            sheet=sheet_name,
                            code_type=code_type,
                            raw_code=normalized,
                            mapped_alias=mapped,
                            context_text=text[:120],
                        )
                    )

    wb.close()
    return extracted


def write_inventory(rows: Sequence[ExtractedCode], output_file: Path) -> None:
    """Write extracted inventory into CSV for documentation workflows."""
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with output_file.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(
            [
                "workbook",
                "sheet",
                "code_type",
                "raw_code",
                "mapped_alias",
                "context_text",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    row.workbook,
                    row.sheet,
                    row.code_type,
                    row.raw_code,
                    row.mapped_alias,
                    row.context_text,
                ]
            )


def main() -> None:
    """Run extraction and persist CSV."""
    index_map = load_index_map(INDEX_MAP_FILE)

    all_rows: List[ExtractedCode] = []
    for workbook in sorted(BASE_XLSX_DIR.glob("*.xlsx")):
        all_rows.extend(scan_workbook(workbook, index_map))

    # Keep stable ordering for easy diff review.
    all_rows.sort(key=lambda r: (r.workbook, r.sheet, r.code_type, r.raw_code))
    write_inventory(all_rows, OUTPUT_FILE)

    print(f"Extracted {len(all_rows)} rows to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
