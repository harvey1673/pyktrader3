"""Backtest every non-zero ``new_weight`` strategy in a weights workbook.

The report uses the same generated historical signals, execution prices,
strategy JSON scalers, btmetrics calculations, and transaction-cost treatment
as :mod:`misc_scripts.strategy_scenario_backtest`.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html as html_lib
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

if __package__ in {None, ""}:
    project_root = str(Path(__file__).resolve().parents[1])
    if project_root not in sys.path:
        sys.path.insert(0, project_root)

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from misc_scripts.strategy_scenario_backtest import (
    BUSINESS_DAYS_PER_YEAR,
    PortfolioResult,
    StrategyScenario,
    _load_saved_daily_futures,
    _safe_ratio,
    _tenor_start,
    btmetrics_result,
    build_generated_historical_provider,
    compose_portfolio,
    load_excel_weight_scenarios,
    load_strategy_scenario,
    trim_portfolio_start,
)


DEFAULT_PORTFOLIO_TENORS = (
    "1m",
    "3m",
    "6m",
    "1y",
    "2y",
    "3y",
    "4y",
    "5y",
    "7y",
    "10y",
    "15y",
)
DEFAULT_SHEET_NAME = "signal_weights"


@dataclass(frozen=True)
class FullPortfolioBacktest:
    """Results and report tables for a workbook-wide proposed portfolio."""

    strategies: Mapping[str, PortfolioResult]
    total: PortfolioResult
    daily_pnl: pd.DataFrame
    total_metrics: pd.DataFrame
    strategy_metrics: pd.DataFrame
    signal_metrics: pd.DataFrame
    tenors: tuple[str, ...]
    start_date: dt.date
    end_date: dt.date
    as_of: dt.date


def workbook_new_weight_strategies(
    weights_excel: str | Path,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> tuple[str, ...]:
    """Return workbook strategies having at least one non-zero new weight."""

    workbook = load_workbook(weights_excel, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Excel sheet '{sheet_name}' not found")
        worksheet = workbook[sheet_name]
        headers: dict[str, int] = {}
        for column, cell in enumerate(worksheet[1], start=1):
            if cell.value is not None:
                headers[str(cell.value).strip().lower().replace(" ", "_")] = column
        missing = [name for name in ("strategy", "new_weight") if name not in headers]
        if missing:
            raise ValueError("Excel sheet is missing columns: " + ", ".join(missing))

        active: dict[str, str] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            strategy_value = row[headers["strategy"] - 1]
            weight_value = row[headers["new_weight"] - 1]
            if strategy_value is None and weight_value is None:
                continue
            if strategy_value is None or not str(strategy_value).strip():
                raise ValueError(f"Row {row_number}: strategy must not be blank")
            try:
                weight = float(weight_value)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Row {row_number}: new_weight must be numeric"
                ) from exc
            if not np.isfinite(weight):
                raise ValueError(f"Row {row_number}: new_weight must be finite")
            if weight == 0.0:
                continue
            strategy = str(strategy_value).strip()
            if Path(strategy).name != strategy:
                raise ValueError(f"Row {row_number}: strategy must be a file name")
            if Path(strategy).suffix.lower() != ".json":
                strategy += ".json"
            active.setdefault(strategy.lower(), strategy)
        return tuple(sorted(active.values(), key=str.lower))
    finally:
        workbook.close()


def _add_frames(frames: Sequence[pd.DataFrame]) -> pd.DataFrame:
    if not frames:
        return pd.DataFrame()
    output = frames[0].copy()
    for frame in frames[1:]:
        output = output.add(frame, fill_value=0.0)
    return output.sort_index().fillna(0.0)


def aggregate_strategy_results(
    results: Mapping[str, PortfolioResult],
) -> PortfolioResult:
    """Combine independently costed production strategies into one portfolio."""

    if not results:
        raise ValueError("No strategy results were supplied")
    values = list(results.values())
    display_keys: dict[tuple[str, str], str] = {}
    used_keys: set[str] = set()
    for strategy, result in results.items():
        strategy_name = Path(strategy).stem
        for factor_name in result.signal_pnl.columns:
            display_key = f"{strategy_name}:{factor_name}"
            if display_key in used_keys:
                raise ValueError(
                    f"Duplicate portfolio signal key '{display_key}'"
                )
            used_keys.add(display_key)
            display_keys[(strategy, factor_name)] = display_key
    signal_pnl = pd.concat(
        [
            result.signal_pnl.rename(
                columns={
                    column: display_keys[(name, column)]
                    for column in result.signal_pnl
                }
            )
            for name, result in results.items()
        ],
        axis=1,
    ).fillna(0.0)
    signal_asset_pnl = {
        display_keys[(name, factor)]: frame
        for name, result in results.items()
        for factor, frame in result.signal_asset_pnl.items()
    }
    signal_holdings = {
        display_keys[(name, factor)]: frame
        for name, result in results.items()
        for factor, frame in result.signal_holdings.items()
    }
    signal_trade_volume = {
        display_keys[(name, factor)]: frame
        for name, result in results.items()
        for factor, frame in result.signal_trade_volume.items()
    }
    signal_gross_exposure = {
        display_keys[(name, factor)]: frame
        for name, result in results.items()
        for factor, frame in result.signal_gross_exposure.items()
    }
    scenario = StrategyScenario(
        name="full_portfolio",
        strategy_file="<weights workbook>",
        scaler=1.0,
        signals={},
        config={},
    )
    return PortfolioResult(
        scenario=scenario,
        cost_mode=values[0].cost_mode,
        signal_source=values[0].signal_source,
        gross_asset_pnl=_add_frames([item.gross_asset_pnl for item in values]),
        costs_by_asset=_add_frames([item.costs_by_asset for item in values]),
        net_asset_pnl=_add_frames([item.net_asset_pnl for item in values]),
        aggregate_holdings=_add_frames([item.aggregate_holdings for item in values]),
        trade_volume=_add_frames([item.trade_volume for item in values]),
        gross_exposure=_add_frames([item.gross_exposure for item in values]),
        signal_pnl=signal_pnl,
        signal_asset_pnl=signal_asset_pnl,
        signal_holdings=signal_holdings,
        signal_trade_volume=signal_trade_volume,
        signal_gross_exposure=signal_gross_exposure,
    )


def _performance_table(
    result: PortfolioResult, tenors: Sequence[str]
) -> pd.DataFrame:
    table = btmetrics_result(result, tenors=tenors).portfolio.rename(
        columns={"std": "daily_std"}
    )
    table.index.name = "tenor"
    return table


def signal_tenor_metrics(
    results: Mapping[str, PortfolioResult],
    tenors: Sequence[str],
    *,
    business_days_per_year: int = BUSINESS_DAYS_PER_YEAR,
) -> pd.DataFrame:
    """Calculate signal Sharpe, risk, turnover, and efficiency by tenor."""

    normalized_tenors = tuple(str(value).strip().lower() for value in tenors)
    rows: list[dict[str, Any]] = []
    for strategy, result in results.items():
        for factor_name, asset_pnl in result.signal_asset_pnl.items():
            pnl = asset_pnl.fillna(0.0).sum(axis=1).sort_index()
            trades = result.signal_trade_volume[factor_name].fillna(0.0)
            exposure = result.signal_gross_exposure[factor_name].fillna(0.0)
            spec = result.scenario.signals[factor_name]
            periods = [("full", pnl.index[0])]
            periods.extend(
                (tenor, _tenor_start(pnl.index[-1], tenor))
                for tenor in normalized_tenors
            )
            for tenor, cutoff in periods:
                sample = pnl.loc[cutoff:]
                sample_trades = trades.loc[cutoff:].sum(axis=1)
                sample_exposure = exposure.loc[cutoff:].sum(axis=1)
                daily_mean = float(sample.mean())
                daily_std = float(sample.std())
                average_trade = float(sample_trades.mean())
                average_exposure = float(sample_exposure.mean())
                cumulative = sample.cumsum()
                drawdown = cumulative - cumulative.cummax()
                max_drawdown = float(drawdown.min()) if len(drawdown) else np.nan
                downside_std = float(sample[sample < 0.0].std())
                rows.append(
                    {
                        "strategy": strategy,
                        "factor_name": factor_name,
                        "signal_name": spec.name,
                        "type": spec.type,
                        "new_weight": spec.weight,
                        "tenor": tenor,
                        "sharpe": _safe_ratio(
                            daily_mean * np.sqrt(business_days_per_year), daily_std
                        ),
                        "daily_std": daily_std,
                        "sortino": _safe_ratio(
                            daily_mean * np.sqrt(business_days_per_year),
                            downside_std,
                        ),
                        "calmar": _safe_ratio(
                            daily_mean * business_days_per_year,
                            -max_drawdown,
                        ),
                        "max_drawdown": max_drawdown,
                        "annualized_pnl": daily_mean * business_days_per_year,
                        "total_pnl": float(sample.sum()),
                        "turnover_pct": 100.0
                        * _safe_ratio(average_trade, average_exposure),
                        "pnl_per_trade_bps": 10000.0
                        * _safe_ratio(daily_mean, average_trade),
                    }
                )
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows).set_index(
        ["strategy", "factor_name", "tenor"]
    ).sort_index()


def run_full_portfolio_backtest(
    settings_dir: str | Path,
    weights_excel: str | Path,
    *,
    start_date: dt.date,
    end_date: dt.date,
    as_of: dt.date | None = None,
    tenors: Sequence[str] = DEFAULT_PORTFOLIO_TENORS,
    sheet_name: str = DEFAULT_SHEET_NAME,
    cost_mode: str = "netted",
    cost_multiplier: float = 1.0,
    holding_lag: int = 2,
    provider_builder: Callable[..., Any] = build_generated_historical_provider,
    progress: Callable[[str], None] | None = print,
) -> FullPortfolioBacktest:
    """Backtest all active workbook strategies using ``new_weight`` only.

    Strategy JSON files remain read-only and supply asset universes, scalers,
    and non-workbook signal settings. Missing JSON files and unresolved signals
    raise errors instead of silently reducing the reported portfolio.
    """

    settings_dir = Path(settings_dir)
    weights_excel = Path(weights_excel)
    if not settings_dir.is_dir():
        raise FileNotFoundError(f"Settings directory does not exist: {settings_dir}")
    if not weights_excel.is_file():
        raise FileNotFoundError(f"Weights workbook does not exist: {weights_excel}")
    if end_date < start_date:
        raise ValueError("end_date must be on or after start_date")
    selected_as_of = end_date if as_of is None else as_of
    if selected_as_of < end_date:
        raise ValueError("as_of must be on or after end_date")
    normalized_tenors = tuple(str(value).strip().lower() for value in tenors)
    strategy_files = workbook_new_weight_strategies(
        weights_excel, sheet_name=sheet_name
    )
    if not strategy_files:
        raise ValueError("The workbook contains no non-zero new_weight strategies")

    results: dict[str, PortfolioResult] = {}
    provider_kwargs: dict[str, Any] = {}
    if provider_builder is build_generated_historical_provider:
        # The production portfolio contains many strategies, but every strategy
        # uses the same as-of futures and fundamental snapshots. Read those large
        # frames once and let each generator work from an isolated copy.
        from misc_scripts.update_fut_prices import load_fun_data

        cached_prices = _load_saved_daily_futures(selected_as_of)
        cached_fundamentals = load_fun_data(selected_as_of)
        provider_kwargs = {
            "price_loader": lambda _as_of: cached_prices,
            "fundamental_loader": lambda _as_of: cached_fundamentals,
        }
    for position, strategy_file in enumerate(strategy_files, start=1):
        if progress:
            progress(f"[{position}/{len(strategy_files)}] Backtesting {strategy_file}")
        template = load_strategy_scenario(
            settings_dir,
            strategy_file,
            scenario_name="template",
        )
        _, proposed = load_excel_weight_scenarios(
            template,
            weights_excel,
            sheet_name=sheet_name,
            proposed_name=Path(strategy_file).stem,
        )
        provider = provider_builder(
            [proposed],
            start_date=start_date,
            end_date=end_date,
            as_of=selected_as_of,
            holding_lag=holding_lag,
            **provider_kwargs,
        )
        result = compose_portfolio(
            proposed,
            provider,
            cost_mode=cost_mode,
            cost_multiplier=cost_multiplier,
        )
        results[strategy_file] = trim_portfolio_start(result, start_date)

    total = aggregate_strategy_results(results)
    daily_pnl = pd.DataFrame(
        {name: result.portfolio_pnl for name, result in results.items()}
    ).fillna(0.0)
    daily_pnl["full_portfolio"] = daily_pnl.sum(axis=1)
    daily_pnl = daily_pnl.sort_index()

    total_metrics = _performance_table(total, normalized_tenors)
    strategy_metrics = pd.concat(
        {
            name: _performance_table(result, normalized_tenors)
            for name, result in results.items()
        },
        names=["strategy", "tenor"],
    )
    signal_metrics = signal_tenor_metrics(results, normalized_tenors)
    return FullPortfolioBacktest(
        strategies=results,
        total=total,
        daily_pnl=daily_pnl,
        total_metrics=total_metrics,
        strategy_metrics=strategy_metrics,
        signal_metrics=signal_metrics,
        tenors=normalized_tenors,
        start_date=start_date,
        end_date=end_date,
        as_of=selected_as_of,
    )


def _format_metrics(frame: pd.DataFrame) -> str:
    display = frame.reset_index().copy()
    for column in ("sharpe", "daily_std"):
        display[column] = display[column].map(
            lambda value: "—" if pd.isna(value) else f"{value:,.4f}"
        )
    return display.to_html(index=False, classes="metrics", border=0, escape=True)


def _format_signal_metrics(frame: pd.DataFrame) -> str:
    display = frame.reset_index().copy()
    numeric_columns = (
        "new_weight",
        "sharpe",
        "daily_std",
        "sortino",
        "calmar",
        "max_drawdown",
        "annualized_pnl",
        "total_pnl",
        "turnover_pct",
        "pnl_per_trade_bps",
    )
    for column in numeric_columns:
        display[column] = display[column].map(
            lambda value: "—" if pd.isna(value) else f"{value:,.4f}"
        )
    display = display.rename(
        columns={
            "turnover_pct": "turnover_%",
            "pnl_per_trade_bps": "pnl_per_trade_bps",
        }
    )
    return display.to_html(index=False, classes="metrics", border=0, escape=True)


def _signal_total_figure(report: FullPortfolioBacktest, go: Any) -> Any:
    signal_pnl = report.total.signal_pnl.fillna(0.0)
    order = signal_pnl.sum().abs().sort_values(ascending=False).index.tolist()
    figure = go.Figure()
    for position, signal in enumerate(order):
        pnl = signal_pnl[signal].cumsum()
        figure.add_trace(
            go.Scatter(
                x=pnl.index,
                y=pnl,
                mode="lines",
                name=signal,
                visible=position == 0,
                line={"width": 2.3, "color": "#245a8d"},
            )
        )
    buttons = []
    for position, signal in enumerate(order):
        visible = [False] * len(order)
        visible[position] = True
        buttons.append(
            {
                "label": signal,
                "method": "update",
                "args": [
                    {"visible": visible},
                    {"title": f"Signal cumulative PNL — {signal}"},
                ],
            }
        )
    title = f"Signal cumulative PNL — {order[0]}" if order else "Signal cumulative PNL"
    figure.update_layout(
        title=title,
        template="plotly_white",
        hovermode="x unified",
        height=530,
        xaxis={"rangeslider": {"visible": True}},
        margin={"l": 65, "r": 25, "t": 115, "b": 90},
        updatemenus=[
            {
                "buttons": buttons,
                "direction": "down",
                "showactive": True,
                "x": 1,
                "y": 1.17,
                "xanchor": "right",
                "yanchor": "top",
            }
        ],
    )
    return figure


def _signal_asset_figure(report: FullPortfolioBacktest, go: Any) -> Any:
    signals = list(report.total.signal_asset_pnl)
    figure = go.Figure()
    groups: dict[str, list[int]] = {}
    palette = (
        "#245a8d",
        "#c67c24",
        "#4f7f52",
        "#a44b62",
        "#7357a3",
        "#327d83",
        "#89613e",
    )
    for signal_position, signal in enumerate(signals):
        asset_pnl = report.total.signal_asset_pnl[signal].fillna(0.0).sort_index()
        assets = [
            str(asset) for asset in asset_pnl if asset_pnl[asset].ne(0.0).any()
        ]
        series = [("Total", asset_pnl.sum(axis=1))]
        series.extend((asset, asset_pnl[asset]) for asset in assets)
        groups[signal] = []
        for series_position, (asset, pnl) in enumerate(series):
            groups[signal].append(len(figure.data))
            figure.add_trace(
                go.Scatter(
                    x=pnl.index,
                    y=pnl.cumsum(),
                    mode="lines",
                    name=asset,
                    visible=signal_position == 0,
                    line={
                        "width": 2.8 if asset == "Total" else 1.4,
                        "dash": "solid" if asset == "Total" else "dot",
                        "color": "#17212b"
                        if asset == "Total"
                        else palette[(series_position - 1) % len(palette)],
                    },
                )
            )
    buttons = []
    for signal in signals:
        visible = [False] * len(figure.data)
        for trace_index in groups[signal]:
            visible[trace_index] = True
        buttons.append(
            {
                "label": signal,
                "method": "update",
                "args": [
                    {"visible": visible},
                    {"title": f"Signal cumulative PNL by asset — {signal}"},
                ],
            }
        )
    title = (
        f"Signal cumulative PNL by asset — {signals[0]}"
        if signals
        else "Signal cumulative PNL by asset"
    )
    figure.update_layout(
        title=title,
        template="plotly_white",
        hovermode="x unified",
        height=570,
        xaxis={"rangeslider": {"visible": True}},
        legend={"orientation": "h", "y": -0.27},
        margin={"l": 65, "r": 25, "t": 115, "b": 120},
        updatemenus=[
            {
                "buttons": buttons,
                "direction": "down",
                "showactive": True,
                "x": 1,
                "y": 1.17,
                "xanchor": "right",
                "yanchor": "top",
            }
        ],
    )
    return figure


def _iplot_figure(frame: pd.DataFrame, title: str, go: Any) -> Any:
    """Build the same one-trace-per-column view as tstool.iplot."""

    figure = go.Figure()
    for column in frame.columns:
        figure.add_trace(
            go.Scatter(x=frame.index, y=frame[column], name=str(column))
        )
    figure.update_layout(
        title=title,
        xaxis_title="date",
        yaxis_title="cumulative PNL",
        width=900,
        height=600,
        template="plotly_white",
        hovermode="x unified",
        dragmode="zoom",
        xaxis={"fixedrange": False, "showgrid": True},
        yaxis={"fixedrange": False, "showgrid": True},
        legend={"orientation": "h", "y": -0.25},
        margin={"l": 65, "r": 25, "t": 65, "b": 100},
    )
    return figure


def _correlation_figure(daily_pnl: pd.DataFrame, go: Any) -> Any:
    """Return a three-year strategy daily-PNL correlation heatmap."""

    if daily_pnl.empty:
        raise ValueError("Strategy daily PNL is empty")
    cutoff = daily_pnl.index[-1] - pd.DateOffset(years=3)
    sample = daily_pnl.loc[cutoff:]
    correlation = sample.corr(min_periods=20)
    figure = go.Figure(
        data=go.Heatmap(
            z=correlation.to_numpy(),
            x=correlation.columns.tolist(),
            y=correlation.index.tolist(),
            zmin=-1.0,
            zmax=1.0,
            zmid=0.0,
            colorscale="RdBu",
            reversescale=True,
            text=correlation.round(2).to_numpy(),
            texttemplate="%{text}",
            hovertemplate="%{y} vs %{x}<br>correlation=%{z:.3f}<extra></extra>",
            colorbar={"title": "Correlation"},
        )
    )
    figure.update_layout(
        title=(
            "Strategy daily-PNL correlation — last 3 years from "
            f"{sample.index[0]:%Y-%m-%d}"
        ),
        template="plotly_white",
        width=1200,
        height=max(700, 32 * len(correlation.index)),
        xaxis={"tickangle": -45, "fixedrange": False},
        yaxis={"autorange": "reversed", "fixedrange": False},
        dragmode="zoom",
        margin={"l": 180, "r": 45, "t": 90, "b": 190},
    )
    return figure


def _format_metric_matrix(frame: pd.DataFrame) -> str:
    display = frame.copy()
    display.index.name = "tenor"
    for column in display.columns:
        display[column] = display[column].map(
            lambda value: "—" if pd.isna(value) else f"{value:,.4f}"
        )
    return display.reset_index().to_html(
        index=False, classes="metrics", border=0, escape=True
    )


def _report_css() -> str:
    return """
body { margin:0; background:#f4f7fa; color:#17212b; font-family:Arial,sans-serif; }
main { max-width:1500px; margin:auto; padding:24px; }
.card { background:white; border:1px solid #dce4ec; border-radius:10px; padding:20px; margin:16px 0; box-shadow:0 2px 7px #13283c12; }
h1,h2,h3 { color:#163a5f; } .note { color:#566575; }
.metrics { border-collapse:collapse; width:100%; font-size:14px; }
.metrics th { background:#163a5f; color:white; text-align:left; padding:8px; position:sticky; top:0; }
.metrics td { border-bottom:1px solid #e2e8ef; padding:7px 8px; }
.metrics tr:nth-child(even) { background:#f7f9fb; }
.scroll { overflow:auto; max-height:620px; }
.wide { overflow:auto; }
.strategy-links { columns:3 280px; padding-left:20px; }
.strategy-links li { margin:8px 0; break-inside:avoid; }
a { color:#245a8d; text-decoration:none; } a:hover { text-decoration:underline; }
.signal { border-top:5px solid #dce4ec; margin-top:34px; padding-top:10px; }
code { background:#eef3f7; padding:2px 5px; border-radius:4px; }
"""


def write_full_portfolio_html(
    report: FullPortfolioBacktest,
    output_html: str | Path,
) -> Path:
    """Write an overview plus one sequential signal report per strategy."""

    try:
        import plotly.graph_objects as go
        from plotly.offline import get_plotlyjs
    except ImportError as exc:
        raise ImportError("HTML report generation requires plotly") from exc

    config = {
        "displaylogo": False,
        "responsive": True,
        "scrollZoom": True,
        "displayModeBar": True,
        "doubleClick": "reset+autosize",
        "modeBarButtonsToAdd": [
            "zoomIn2d",
            "zoomOut2d",
            "autoScale2d",
            "resetScale2d",
        ],
    }
    output_html = Path(output_html)
    output_html.parent.mkdir(parents=True, exist_ok=True)
    strategy_dir = output_html.parent / f"{output_html.stem}_strategies"
    asset_dir = output_html.parent / f"{output_html.stem}_assets"
    strategy_dir.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)
    plotly_asset = asset_dir / "plotly.min.js"
    plotly_asset.write_text(get_plotlyjs(), encoding="utf-8")

    cumulative = report.daily_pnl.cumsum()
    total_figure = _iplot_figure(
        cumulative[["full_portfolio"]].rename(
            columns={"full_portfolio": "Full portfolio"}
        ),
        "Full portfolio cumulative PNL",
        go,
    )
    strategy_figure = _iplot_figure(
        cumulative.loc[:, list(report.strategies)],
        "Cumulative PNL by strategy — click legend entries to select/unselect",
        go,
    )
    strategy_daily_pnl = report.daily_pnl.loc[:, list(report.strategies)]
    correlation_figure = _correlation_figure(strategy_daily_pnl, go)

    page_names: dict[str, str] = {}
    for strategy in report.strategies:
        safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(strategy).stem)
        page_name = f"{safe_name}.html"
        if page_name in page_names.values():
            raise ValueError(f"Duplicate strategy report filename '{page_name}'")
        page_names[strategy] = page_name

    links = "\n".join(
        f'<li><a href="{html_lib.escape(strategy_dir.name)}/{html_lib.escape(page_name)}">'
        f"{html_lib.escape(strategy)}</a></li>"
        for strategy, page_name in page_names.items()
    )
    tenor_order = ["full", *report.tenors]
    strategy_metric_rows = report.strategy_metrics.reset_index()
    sharpe_matrix = strategy_metric_rows.pivot(
        index="tenor", columns="strategy", values="sharpe"
    ).reindex(index=tenor_order, columns=list(report.strategies))
    std_matrix = strategy_metric_rows.pivot(
        index="tenor", columns="strategy", values="daily_std"
    ).reindex(index=tenor_order, columns=list(report.strategies))
    overview = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Production portfolio weight backtest</title>
<script src="{html_lib.escape(asset_dir.name)}/plotly.min.js"></script>
<style>{_report_css()}</style></head><body><main>
<h1>Production portfolio weight backtest</h1>
<p class="note">Workbook proposed weights (<code>new_weight</code>) · {report.start_date:%Y-%m-%d} to {report.end_date:%Y-%m-%d} · data as of {report.as_of:%Y-%m-%d} · {len(report.strategies)} strategies · {len(report.total.signal_pnl.columns)} active signals · metrics annualized with {BUSINESS_DAYS_PER_YEAR} business days. Drag a box inside any chart to zoom both axes; use the modebar or double-click to reset.</p>
<section class="card">{total_figure.to_html(full_html=False, include_plotlyjs=False, config=config)}</section>
<section class="card"><h2>Full portfolio performance</h2>{_format_metrics(report.total_metrics.reindex(tenor_order))}</section>
<section class="card">{strategy_figure.to_html(full_html=False, include_plotlyjs=False, config=config)}</section>
<section class="card">{correlation_figure.to_html(full_html=False, include_plotlyjs=False, config=config)}</section>
<section class="card"><h2>Strategy reports</h2><ul class="strategy-links">{links}</ul></section>
<section class="card"><h2>Strategy Sharpe by tenor</h2><div class="wide">{_format_metric_matrix(sharpe_matrix)}</div></section>
<section class="card"><h2>Strategy daily standard deviation by tenor</h2><div class="wide">{_format_metric_matrix(std_matrix)}</div></section>
</main></body></html>"""
    output_html.write_text(overview, encoding="utf-8")

    for strategy, result in report.strategies.items():
        strategy_metric = report.strategy_metrics.xs(
            strategy, level="strategy"
        ).reindex(tenor_order)
        portfolio_curve = _iplot_figure(
            result.portfolio_pnl.cumsum().to_frame("total"),
            f"{strategy} cumulative PNL",
            go,
        )
        signal_sections: list[str] = []
        for factor_name, asset_pnl in result.signal_asset_pnl.items():
            spec = result.scenario.signals[factor_name]
            total_curve = asset_pnl.fillna(0.0).sum(axis=1).cumsum().to_frame("total")
            active_assets = [
                asset for asset in asset_pnl if asset_pnl[asset].fillna(0.0).ne(0.0).any()
            ]
            asset_curve = asset_pnl.loc[:, active_assets].fillna(0.0).cumsum()
            total_chart = _iplot_figure(
                total_curve,
                f"{factor_name} cumulative PNL",
                go,
            )
            asset_chart = _iplot_figure(
                asset_curve,
                f"{factor_name} cumulative PNL by asset",
                go,
            )
            metric = report.signal_metrics.xs(
                (strategy, factor_name), level=("strategy", "factor_name")
            ).reindex(tenor_order)
            signal_sections.append(
                f"""<section class="card signal">
<h2>{html_lib.escape(factor_name)}</h2>
<p class="note">Signal: <code>{html_lib.escape(spec.name)}</code> · type: <code>{html_lib.escape(spec.type)}</code> · new weight: {spec.weight:,.6g}</p>
{total_chart.to_html(full_html=False, include_plotlyjs=False, config=config)}
{asset_chart.to_html(full_html=False, include_plotlyjs=False, config=config)}
<h3>Performance by tenor</h3><div class="scroll">{_format_signal_metrics(metric)}</div>
</section>"""
            )
        strategy_page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html_lib.escape(strategy)} backtest</title>
<script src="../{html_lib.escape(asset_dir.name)}/plotly.min.js"></script>
<style>{_report_css()}</style></head><body><main>
<p><a href="../{html_lib.escape(output_html.name)}">← Portfolio overview</a></p>
<h1>{html_lib.escape(strategy)}</h1>
<p class="note">{report.start_date:%Y-%m-%d} to {report.end_date:%Y-%m-%d} · data as of {report.as_of:%Y-%m-%d} · scaler {result.scenario.scaler:,.6g} · {len(result.signal_asset_pnl)} active signals.</p>
<section class="card">{portfolio_curve.to_html(full_html=False, include_plotlyjs=False, config=config)}</section>
<section class="card"><h2>Strategy performance by tenor</h2>{_format_metrics(strategy_metric)}</section>
{''.join(signal_sections)}
</main></body></html>"""
        (strategy_dir / page_names[strategy]).write_text(
            strategy_page, encoding="utf-8"
        )
    return output_html.resolve()


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("settings_dir", type=Path)
    parser.add_argument("weights_excel", type=Path)
    parser.add_argument("output_html", type=Path)
    parser.add_argument("--start-date", type=dt.date.fromisoformat, required=True)
    parser.add_argument("--end-date", type=dt.date.fromisoformat, required=True)
    parser.add_argument("--as-of", type=dt.date.fromisoformat)
    parser.add_argument("--sheet-name", default=DEFAULT_SHEET_NAME)
    parser.add_argument("--cost-mode", choices=("netted", "sleeve"), default="netted")
    parser.add_argument("--cost-multiplier", type=float, default=1.0)
    parser.add_argument("--holding-lag", type=int, default=2)
    parser.add_argument("--tenors", nargs="+", default=list(DEFAULT_PORTFOLIO_TENORS))
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    report = run_full_portfolio_backtest(
        args.settings_dir,
        args.weights_excel,
        start_date=args.start_date,
        end_date=args.end_date,
        as_of=args.as_of,
        tenors=args.tenors,
        sheet_name=args.sheet_name,
        cost_mode=args.cost_mode,
        cost_multiplier=args.cost_multiplier,
        holding_lag=args.holding_lag,
    )
    output = write_full_portfolio_html(report, args.output_html)
    print(f"Wrote portfolio overview to {output}")
    print(
        "Wrote individual strategy reports to "
        f"{output.parent / (output.stem + '_strategies')}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
