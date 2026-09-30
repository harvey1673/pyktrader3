"""Read-only workbook inventory for the physical-carry research shortlist.

Uses cached Excel values, never refreshes or modifies source workbooks.
Run with --source-dir PATH; outputs are written beside this script by default.
Coverage denominators are dated worksheet rows, not exchange trading calendars.
"""
from __future__ import annotations

import argparse
import ast
import csv
import datetime as dt
import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def load_aliases(root):
    tree = ast.parse((root / 'tests/index_map_full.py').read_text(encoding='utf-8'))
    return next(ast.literal_eval(n.value) for n in tree.body
                if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name)
                and t.id == 'index_map_full' for t in n.targets))


def product_for(alias, description):
    if '电价' in description:
        return None
    for prefix, product in [('coke_', 'J'), ('ckc_', 'JM'), ('ss_304', 'SS'),
                            ('sm_65', 'SM'), ('sf_72', 'SF'), ('sf_75', 'SF')]:
        if alias.startswith(prefix) and not any(x in alias for x in ['inv', 'cost', 'profit']):
            return product
    if any(x in description for x in ['价', '基差', '升贴水']):
        for keyword, product in [('焦煤', 'JM'), ('主焦煤', 'JM'), ('焦炭', 'J'),
                                 ('硅锰', 'SM'), ('锰硅', 'SM'), ('硅铁', 'SF'),
                                 ('不锈钢', 'SS')]:
            if keyword in description:
                return product
    return None


def iso(value):
    return value.isoformat() if isinstance(value, (dt.datetime, dt.date)) else value


def write_csv(path, rows):
    if not rows:
        return
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, default=Path(__file__).parent / 'priority_spot_audit')
    parser.add_argument('--as-of', type=dt.date.fromisoformat, default=dt.date.today())
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[2]
    aliases = load_aliases(root)
    sources = {'ifind_daily_full.xlsx': ['ferrous_d', 'base_d'],
               'ifind_daily.xlsx': ['ferrous_d', 'base_d'],
               'ifind_data.xlsx': ['hist', 'const_d', 'base_d2', 'ferrous_w'],
               'mysteel_metal.xlsx': ['data']}
    inventory, yearly, observations, scanned = [], [], [], []
    for filename, sheets in sources.items():
        path = args.source_dir / filename
        if not path.exists():
            scanned.append({'file': filename, 'status': 'missing'})
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in sheets:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            it = ws.iter_rows(values_only=True)
            headers = [next(it, ()) for _ in range(10)]
            # iFinD metadata: description/frequency/unit/id in Excel rows 3:6.
            # Mysteel metadata: description/unit/id/frequency in rows 2,3,5,6.
            name_r, code_r, freq_r, unit_r = (1, 4, 5, 2) if filename.startswith('mysteel') else (2, 5, 3, 4)
            if sheet == 'hist':
                name_r, code_r, freq_r, unit_r = 3, 6, 4, 5
            selected = {}
            for col, raw_code in enumerate(headers[code_r]):
                code = str(raw_code or '').strip().split('.')[0]
                alias = aliases.get(code, '')
                description = str(headers[name_r][col] or '')
                product = product_for(alias, description)
                if product:
                    selected[col] = {'product': product, 'code': code, 'alias': alias,
                        'description': description, 'unit': headers[unit_r][col],
                        'frequency': headers[freq_r][col], 'workbook': filename,
                        'sheet': sheet, 'column': get_column_letter(col + 1),
                        'header_range': f'{get_column_letter(col+1)}{name_r+1}:{get_column_letter(col+1)}9',
                        'advertised_range': str(headers[6][col] or '') if filename.startswith('ifind') else '',
                        'header_update': iso(headers[8][col]) if filename.startswith('ifind') else ''}
            collected = defaultdict(list)
            dates = []
            # Row 10 can already contain data; include it.
            def data_rows():
                yield 10, headers[9]
                yield from enumerate(it, 11)
            for row_number, row in data_rows():
                if not row or not isinstance(row[0], (dt.datetime, dt.date)):
                    continue
                date = row[0].date() if isinstance(row[0], dt.datetime) else row[0]
                if date < dt.date(2016, 1, 1) or date > args.as_of:
                    continue
                dates.append(date)
                for col, meta in selected.items():
                    value = row[col] if col < len(row) else None
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
                        collected[col].append((date, float(value), row_number))
                    elif value is not None:
                        collected[(col, 'nonnumeric')].append(str(value))
            scanned.append({'file': filename, 'sheet': sheet, 'worksheet_rows': ws.max_row,
                            'dated_rows_since_2016': len(dates), 'selected_columns': len(selected)})
            for col, meta in selected.items():
                data = sorted(collected[col])
                valid_dates = sorted(set(x[0] for x in data))
                positive = [(d, v, r) for d, v, r in data if v > 0]
                duplicate_count = len(data) - len(valid_dates)
                by_date = defaultdict(set)
                for d, v, r in data:
                    by_date[d].add(v)
                    observations.append({**{k: meta[k] for k in ['product','code','alias','workbook','sheet']},
                                         'date': d.isoformat(), 'value': v,
                                         'cell': f'{meta["column"]}{r}'})
                inventory.append({**meta, 'numeric_count': len(data),
                    'positive_count': len(positive), 'nonpositive_count': len(data)-len(positive),
                    'nonnumeric_nonempty_count': len(collected[(col, 'nonnumeric')]),
                    'first_numeric': valid_dates[0].isoformat() if valid_dates else '',
                    'last_numeric': valid_dates[-1].isoformat() if valid_dates else '',
                    'first_positive': positive[0][0].isoformat() if positive else '',
                    'last_positive': positive[-1][0].isoformat() if positive else '',
                    'max_gap_calendar_days': max(((b-a).days for a,b in zip(valid_dates,valid_dates[1:])), default=0),
                    'duplicate_dates': duplicate_count,
                    'conflicting_duplicate_dates': sum(len(v)>1 for v in by_date.values()),
                    'latest_value': data[-1][1] if data else None})
                for year in sorted(set(d.year for d in dates)):
                    yd = [v for d,v,r in data if d.year == year]
                    yearly.append({**{k: meta[k] for k in ['product','code','alias','workbook','sheet','column']},
                        'year': year, 'dated_sheet_rows': sum(d.year == year for d in dates),
                        'numeric_count': len(yd), 'positive_count': sum(v>0 for v in yd)})
            print(f'{filename}/{sheet}: {len(selected)} candidate columns, {len(dates)} dated rows', flush=True)
        wb.close()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / 'inventory.csv', inventory)
    write_csv(args.output_dir / 'yearly_coverage.csv', yearly)
    # Preserve source-specific values; never silently merge conflicting workbook snapshots.
    write_csv(args.output_dir / 'observations.csv', observations)
    grouped = defaultdict(list)
    for r in observations:
        grouped[(r['code'],r['date'])].append(r)
    conflicts = []
    for (code,date), rows in grouped.items():
        if len(set(r['value'] for r in rows)) > 1:
            conflicts.append({'code': code, 'date': date,
                'source_values': json.dumps([{k:r[k] for k in ['workbook','sheet','cell','value']} for r in rows])})
    write_csv(args.output_dir / 'overlap_conflicts.csv', conflicts)
    summary = {'source_directory': str(args.source_dir), 'start_date': '2016-01-01', 'as_of': args.as_of.isoformat(),
        'extraction': 'read_only cached numeric values; no refresh or forward-fill',
        'coverage_denominator': 'dated worksheet rows, not trading days',
        'scanned': scanned, 'inventory_rows': len(inventory), 'observation_rows': len(observations),
        'conflicting_code_dates': len(conflicts)}
    (args.output_dir / 'audit_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=True), flush=True)


if __name__ == '__main__':
    main()
