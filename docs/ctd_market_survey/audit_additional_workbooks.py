"""Read-only metadata scan and block-local Mysteel date audit."""
import datetime as dt
import json
from pathlib import Path
import openpyxl

ROOT = Path('C:/Users/harve/Nutstore/1/Nutstore')
OUT = Path(__file__).parent / 'priority_spot_audit'
headers, profiles = [], []
keywords = ['焦煤', '焦炭', '冶金焦', '硅铁', '硅锰', '锰硅', '304', '蒙5', '蒙五']
for path in ROOT.glob('*.xlsx'):
    if not path.name.startswith(('ifind', 'mysteel')):
        continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb:
        top = list(ws.iter_rows(max_row=15, values_only=True))
        width = max(map(len, top), default=0)
        for col in range(1, width):
            vals = []
            for row in top:
                value = row[col] if col < len(row) else None
                if isinstance(value, (dt.datetime, dt.date, int, float)):
                    break
                if value is not None:
                    vals.append(str(value))
            description = ' | '.join(vals)
            if any(k in description for k in keywords):
                headers.append(dict(workbook=path.name, sheet=ws.title,
                    column=openpyxl.utils.get_column_letter(col+1), metadata=description))
        if not path.name.startswith('mysteel data') or ws.title != 'mysteel prices':
            continue
        columns = ['N', 'AF', 'DT', 'EE', 'EN', 'FT', 'FU']
        selected = {}
        for column in columns:
            ci = openpyxl.utils.column_index_from_string(column)-1
            # Every export block has its own date column. Column A is NOT a shared index.
            di = max(i for i in range(ci) if top[0][i] == '钢联数据')
            start = next((i for i,r in enumerate(top) if isinstance(r[di], dt.datetime)), len(top))
            selected[column] = (ci, di, start, [])
        for row in ws.iter_rows(values_only=True):
            for ci, di, start, obs in selected.values():
                date, value = row[di], row[ci]
                if isinstance(date, dt.datetime) and isinstance(value, (int,float)) and 2016 <= date.year <= 2026:
                    obs.append((date.date().isoformat(), value))
        for column, (ci,di,start,obs) in selected.items():
            profiles.append(dict(workbook=path.name, sheet=ws.title,column=column,
                date_column=openpyxl.utils.get_column_letter(di+1),
                metadata=[str(r[ci]) for r in top[:start]],count=len(obs),
                first=min((d for d,v in obs),default=None),last=max((d for d,v in obs),default=None)))
    wb.close()
for name, rows in [('all_workbook_headers.json',headers),('additional_mysteel_profiles.json',profiles)]:
    (OUT/name).write_text(json.dumps(rows, ensure_ascii=False, indent=2),encoding='utf-8')
for r in profiles:
    print(r['workbook'],r['column'],r['date_column'],r['count'],r['first'],r['last'])
